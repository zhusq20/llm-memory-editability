#!/usr/bin/env python3
"""Phase 6 (Reporter): build the screening deliverables from merged decisions.

Usage:
  python build_outputs.py --work W --out OUT_DIR [--stage ta|ft] [--config screening_config.json]
                          [--tag-keywords]

Writes into OUT_DIR (prefix TA_ or FT_):
  <P>_screening_log.xlsx     Summary, advanced/included records, all records, conflicts, QC, pending
                             (CSV files instead when openpyxl is not installed)
  <P>_1_include.ris, <P>_2_unclear.ris, <P>_3_exclude.ris
                             for EndNote / Zotero / Mendeley: Label (LB) = decision, note (N1) = reasons;
                             --tag-keywords also adds the decision as a keyword (Zotero tag)
  <P>_prisma_counts.json / .md   PRISMA 2020 numbers (+ a Mermaid flow diagram)
  <P>_methods_selection.md       methods paragraph and AI-use statement, with [TO COMPLETE] slots
  <P>_literature_corpus.yaml     advanced (TA) or included (FT) records as literature_corpus[]
                                 entries for Academic Research Skills (academic-paper / deep-research)
Nothing here is final while records are pending: the summary says so and no methods text is written.
"""
import argparse
import csv
import datetime
import json
import os
import re
import sys
import unicodedata

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import srlib  # noqa: E402

REPO_URL = "https://github.com/Imbad0202/academic-research-skills"


def split_author(a):
    a = a.strip().rstrip(".")
    if not a:
        return None
    if "," in a:
        fam, giv = a.split(",", 1)
        if not fam.strip():
            return {"literal": a}
        return {"family": fam.strip(), "given": giv.strip()} if giv.strip() else {"family": fam.strip()}
    parts = a.split()
    if len(parts) >= 2 and re.fullmatch(r"([A-Z]\.?-?){1,4}", parts[-1]):
        return {"family": " ".join(parts[:-1]), "given": parts[-1]}
    return {"literal": a}


def ascii_key(s):
    s = unicodedata.normalize("NFKD", s or "")
    return re.sub(r"[^a-z0-9]", "", s.encode("ascii", "ignore").decode().lower())


def yq(v):
    return json.dumps(v, ensure_ascii=False)


def write_corpus(path, rows, stage, when):
    """ARS literature_corpus[] entries (shared/contracts/passport/literature_corpus_entry.schema.json)."""
    ok, rejected = [], []
    for r in rows:
        u = r["u"]
        authors = [x for x in (split_author(a) for a in u.get("authors", [])) if x]
        if not authors:
            rejected.append((u["id"], "no authors"))
            continue
        if not str(u.get("year", "")).isdigit():
            rejected.append((u["id"], "no publication year"))
            continue
        first = authors[0]
        fam = ascii_key(first.get("family") or first.get("literal", ""))[:20] or "rec"
        if not fam[0].isalpha():
            fam = "rec" + fam
        pointer = (f"https://doi.org/{u['doi']}" if u.get("doi") else
                   f"https://pubmed.ncbi.nlm.nih.gov/{u['pmid']}/" if u.get("pmid") else f"urn:sr-screener:{u['id']}")
        e = [("citation_key", f"{fam}{u['year']}_{u['id']}"), ("title", u["title"] or "[no title]"),
             ("authors", authors), ("year", int(u["year"])), ("source_pointer", pointer)]
        if u.get("journal"):
            e.append(("venue", u["journal"]))
        if u.get("doi") and srlib.DOI_RE.match(u["doi"]):
            e.append(("doi", u["doi"]))
        e += [("tags", [srlib.TOOL, f"{stage.upper()}-{r['f']['d'].capitalize()}"]),
              ("obtained_via", "other"), ("adapter_name", srlib.TOOL), ("adapter_version", srlib.VERSION),
              ("obtained_at", when),
              ("user_notes", f"{stage.upper()} screening: {r['f']['d']} ({r['f']['code']}) - {r['f']['why']} "
                             f"[decided by {r['f']['by']}]")]
        ok.append(e)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("# literature_corpus[] entries for Academic Research Skills (literature_corpus_entry.schema.json).\n")
        f.write(f"# Produced by {srlib.TOOL} {srlib.VERSION}. Paste under `literature_corpus:` in a Material Passport,\n")
        f.write("# or give the file to academic-paper / deep-research as your curated corpus.\n")
        if not ok:
            f.write("[]\n")
        for e in ok:
            lead = "- "
            for k, v in e:
                if k == "authors":
                    f.write(f"{lead}authors:\n")
                    for au in v:
                        items = list(au.items())
                        f.write(f"    - {items[0][0]}: {yq(items[0][1])}\n")
                        for kk, vv in items[1:]:
                            f.write(f"      {kk}: {yq(vv)}\n")
                elif k == "tags":
                    f.write(f"{lead}tags:\n")
                    for t in v:
                        f.write(f"    - {yq(t)}\n")
                else:
                    f.write(f"{lead}{k}: {yq(v)}\n")
                lead = "  "
    return len(ok), rejected


def lang_ok(u, allowed):
    if not allowed:
        return True
    langs = [x.strip().lower() for x in (u.get("lang") or "").split(";") if x.strip()]
    return not langs or any(x in [a.lower() for a in allowed] for x in langs)


def main():
    srlib.utf8_stdout()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--work", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--stage", choices=["ta", "ft"], default="ta")
    ap.add_argument("--config")
    ap.add_argument("--tag-keywords", action="store_true")
    a = ap.parse_args()

    cfg = srlib.load_config(a.config)
    work, out = os.path.abspath(a.work), os.path.abspath(a.out)
    os.makedirs(out, exist_ok=True)
    P = a.stage.upper()
    labels = srlib.code_labels(cfg, a.stage)
    recs = srlib.load_json(os.path.join(work, "records.json"))
    if recs is None:
        raise SystemExit("records.json not found")
    U = {u["id"]: u for u in recs["unique"]}
    RAW = srlib.load_json(os.path.join(work, "raw_records.json"), [])
    ident = srlib.load_json(os.path.join(work, "identification.json"), {})
    ta_dec = srlib.load_json(os.path.join(work, "decisions.json"), {})
    ta_pending_ids, ta_qc_ids, state_problems = srlib.ta_status(work, cfg)
    ft_stale = False
    if a.stage == "ta":
        D = ta_dec
        agree = srlib.load_json(os.path.join(work, "agreement.json"), {})
        pending_ids = sorted(ta_pending_ids)
        universe = [u["id"] for u in recs["unique"]]
    else:
        D = srlib.load_json(os.path.join(work, "ft_decisions.json"), {})
        pending = srlib.load_json(os.path.join(work, "ft_pending.json"), {"items": [], "adj": []})
        agree = srlib.load_json(os.path.join(work, "ft_agreement.json"), {})
        man = srlib.load_json(os.path.join(work, "ft_manifest.json"), [])
        universe = [x["id"] for x in man]
        ta_pending_ids, ta_qc_ids, state_problems, uncovered, ft_stale = srlib.fulltext_status(work, cfg)
        missing = {rid for rid in universe if not srlib.valid_decision(D.get(rid, {}).get("final"),
                                                                       srlib.exclusion_codes(cfg, "ft"))}
        if universe and not srlib.decision_state_current(work, "ft", cfg):
            state_problems.append("full-text decisions have no current protocol/config/retrieval identity")
            missing |= set(universe)
        pending_ids = sorted({x["id"] for x in pending["items"]} | {x["id"] for x in pending["adj"]} |
                             ta_pending_ids | missing | uncovered)

    rows = []
    for rid in universe:
        r = D.get(rid)
        if not r or rid not in U or not srlib.valid_decision(r.get("final"), srlib.exclusion_codes(cfg, a.stage)):
            continue
        u = U[rid]
        rows.append(dict(u=u, f=r["final"], A=r.get("A") or {}, B=r.get("B") or {}, J=r.get("ADJ") or {},
                         Q=r.get("QC") or {}, pre_qc=r.get("pre_qc"), pre_override=r.get("pre_override"),
                         qc_flag=r.get("qc_flag", False), ok_lang=lang_ok(u, cfg.get("languages_allowed"))))
    complete = not pending_ids and not state_problems and not ft_stale
    count = lambda d: sum(1 for r in rows if r["f"]["d"] == d)
    exc_codes = {}
    for r in rows:
        if r["f"]["d"] == "exclude":
            exc_codes[r["f"]["code"]] = exc_codes.get(r["f"]["code"], 0) + 1
    code_order = srlib.exclusion_codes(cfg, a.stage)
    exc_sorted = sorted(exc_codes.items(), key=lambda kv: code_order.index(kv[0]) if kv[0] in code_order else 99)
    adv_rows = [r for r in rows if r["f"]["d"] in srlib.ADVANCE] if a.stage == "ta" else \
        [r for r in rows if r["f"]["d"] == "include"]
    seeds = srlib.load_json(os.path.join(work, "seeds.json"), []) or []
    by = {}
    for r in rows:
        by[r["f"]["by"]] = by.get(r["f"]["by"], 0) + 1
    models = cfg["models"]
    mlab = lambda k: cfg.get("model_labels", {}).get(k) or f"[model for {k}: {models.get(k) or 'session model'}]"

    def pair(k1, k2):
        if mlab(k1) == mlab(k2):
            return f"both {mlab(k1)}, prompted with different expert personas"
        return f"{mlab(k1)}; {mlab(k2)}"

    # --------------------------------------------------------------- counts
    counts = {"stage": a.stage, "tool": f"{srlib.TOOL} {srlib.VERSION}", "complete": complete,
              "pending_records": len(pending_ids)}
    if a.stage == "ft":
        counts.update({"pending_ta_records": len(ta_pending_ids), "pending_ta_qc_records": len(ta_qc_ids),
                       "fulltext_set_out_of_date": ft_stale})
    counts["state_problems"] = state_problems
    if a.stage == "ta":
        counts.update({
            "identified_by_database": ident.get("by_database", {}),
            "identified_total": ident.get("raw_total"),
            "duplicates_removed": ident.get("duplicates_removed"),
            "records_screened": len(rows),
            "records_to_screen": len(universe),
            "records_excluded": count("exclude"),
            "excluded_by_code": dict(exc_sorted),
            "reports_sought_for_retrieval": count("include") + count("unclear"),
            "advanced_include": count("include"),
            "advanced_unclear": count("unclear"),
            "advanced_outside_allowed_languages": sum(1 for r in adv_rows if not r["ok_lang"]),
        })
        srlib.save_json(os.path.join(work, "prisma_ta.json"), counts, indent=1)
    else:
        ta_counts = srlib.load_json(os.path.join(work, "prisma_ta.json"), {})
        nr = srlib.load_json(os.path.join(work, "ft_not_retrieved.json"), []) or []
        counts.update({
            "reports_sought_for_retrieval": sum(1 for r in ta_dec.values()
                                                if r.get("final", {}).get("d") in srlib.ADVANCE),
            "reports_not_retrieved": len(nr),
            "reports_assessed": len(rows),
            "reports_to_assess": len(universe),
            "reports_excluded": count("exclude"),
            "excluded_by_reason": {labels.get(k, k): v for k, v in exc_sorted},
            "reports_unclear_awaiting_classification": count("unclear"),
            "studies_included_reports": count("include"),
            "ta": ta_counts,
        })
    counts["agreement"] = agree
    counts["decided_by"] = by
    srlib.save_json(os.path.join(out, f"{P}_prisma_counts.json"), counts, indent=1)

    # --------------------------------------------------------------- sheets
    head = ["ID", "Final decision", "Final code", "Final reason", "Decided by", "Reviewer A", "A code", "A reason",
            "Reviewer B", "B code", "B reason", "Adjudicator", "ADJ reason", "QC", "QC reason", "Before QC/override",
            "Language OK", "First author", "Year", "Title", "Journal", "Publication type", "Language", "DOI", "PMID",
            "Sources", "Batch"]
    if a.stage == "ft":
        head[3:4] = ["Final reason", "Where"]
    head.append("Abstract")

    def line(r):
        u, f = r["u"], r["f"]
        before = r["pre_override"] or r["pre_qc"]
        row = [u["id"], f["d"], f["code"], f.get("why", "")]
        if a.stage == "ft":
            row.append(f.get("where", ""))
        row += [f["by"], r["A"].get("d", ""), r["A"].get("code", ""), r["A"].get("why", ""),
                r["B"].get("d", ""), r["B"].get("code", ""), r["B"].get("why", ""),
                (r["J"].get("d", "") + "/" + r["J"].get("code", "")) if r["J"] else "", r["J"].get("why", ""),
                (r["Q"].get("d", "") + "/" + r["Q"].get("code", "")) if r["Q"] else ("FLAG" if r["qc_flag"] else ""),
                r["Q"].get("why", ""), (before["d"] + "/" + before["code"]) if before else "",
                "yes" if r["ok_lang"] else "NO", (u.get("authors") or [""])[0], u["year"], u["title"], u["journal"],
                u["type"], u["lang"], u["doi"], u["pmid"], ", ".join(u["sources"]), u.get("batch", ""),
                (u["abstract"] or "")[:32000]]
        return row

    summary = [("sr-screener " + srlib.VERSION + f" - {'title/abstract' if a.stage == 'ta' else 'full-text'} "
                "screening log", "")]
    if not complete:
        summary.append((f"INCOMPLETE: {len(pending_ids)} records need decisions or required rechecks "
                        "(resume with build_workflow.py --jobs pending, then --jobs recheck for the required QC "
                        "recheck). Numbers below are provisional.", ""))
        summary += [("INCOMPLETE: " + problem, "") for problem in state_problems]
        if a.stage == "ft" and ta_pending_ids:
            summary.append((f"Title/abstract stage still pending: {len(ta_pending_ids)} records, "
                            f"including {len(ta_qc_ids)} required title/abstract QC rechecks. "
                            "Full-text counts cannot be final until these are decided.", ""))
    if a.stage == "ta":
        summary += [("IDENTIFICATION", "")]
        summary += [(f"Records from {db}", n) for db, n in counts["identified_by_database"].items()]
        summary += [("Total records identified", counts["identified_total"]),
                    ("Duplicate records removed", counts["duplicates_removed"]),
                    ("SCREENING (title/abstract)", ""),
                    ("Records screened", counts["records_screened"]),
                    ("Records with a final decision", len(rows)),
                    ("Records excluded", counts["records_excluded"])]
        summary += [("   " + labels.get(c, c), n) for c, n in exc_sorted]
        summary += [("Records advanced to full-text assessment", counts["reports_sought_for_retrieval"]),
                    ("   include (eligible on title/abstract)", counts["advanced_include"]),
                    ("   unclear (needs the full text)", counts["advanced_unclear"])]
        if cfg.get("languages_allowed"):
            summary.append((f"   of which language outside {', '.join(cfg['languages_allowed'])} (flag, not excluded)",
                            counts["advanced_outside_allowed_languages"]))
    else:
        summary += [("FULL-TEXT ASSESSMENT", ""),
                    ("Reports sought for retrieval", counts["reports_sought_for_retrieval"]),
                    ("Reports not retrieved", counts["reports_not_retrieved"]),
                    ("Reports assessed for eligibility", counts["reports_assessed"]),
                    ("Reports excluded", counts["reports_excluded"])]
        summary += [("   " + k, v) for k, v in counts["excluded_by_reason"].items()]
        summary += [("Awaiting classification (unclear)", counts["reports_unclear_awaiting_classification"]),
                    ("Reports of included studies", counts["studies_included_reports"])]
    summary += [("REVIEWER AGREEMENT", ""),
                ("Records decided by both reviewers", agree.get("n")),
                ("Observed agreement", agree.get("observed_agreement")),
                ("Cohen's kappa", agree.get("kappa")),
                ("PABAK (prevalence-adjusted)", agree.get("pabak")),
                ("Conflicts", agree.get("conflicts", agree.get("a_only_advance", 0) + agree.get("b_only_advance", 0)))]
    summary += [(f"Final decisions made by {k}", v) for k, v in sorted(by.items())]
    if seeds:
        summary.append(("SEED STUDIES (sensitivity check)", ""))
        for s in seeds:
            fd = D.get(s["id"], {}).get("final", {}) if s.get("id") else {}
            summary.append((f"{s['label']}", f"{s['id']}: {fd.get('d', 'pending')}" if s.get("id") else "NOT IN SEARCH"))
    summary += [("", ""), ("AI decisions are decision support. The review team must verify them before reporting "
                           "(see the methods file for the reporting text).", "")]

    cols = [("Advanced" if a.stage == "ta" else "Included", adv_rows),
            ("All_Records", rows),
            ("Conflicts", [r for r in rows if r["f"]["by"] in ("ADJ", "LIBERAL")]),
            ("QC_and_Overrides", [r for r in rows if r["Q"] or r["qc_flag"] or r["pre_override"]])]
    if a.stage == "ft":
        cols.insert(1, ("Excluded_with_reasons", [r for r in rows if r["f"]["d"] == "exclude"]))
    order = {"include": 0, "unclear": 1, "exclude": 2}
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws.column_dimensions["A"].width = 78
        ws.column_dimensions["B"].width = 18
        for k, v in summary:
            ws.append([srlib.spreadsheet_text(k), srlib.spreadsheet_text(v)])
            if k and k.isupper():
                ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        fills = {"include": "C6EFCE", "unclear": "FFEB9C", "exclude": "F2F2F2"}
        widths = {"ID": 9, "Title": 60, "Abstract": 80, "Final reason": 40, "A reason": 32, "B reason": 32,
                  "ADJ reason": 32, "QC reason": 32, "Journal": 28, "Publication type": 24, "Sources": 20, "DOI": 24}
        for name, rs in cols:
            sh = wb.create_sheet(name)
            sh.append(head)
            for i, h in enumerate(head, 1):
                c = sh.cell(row=1, column=i)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E78")
                c.alignment = Alignment(wrap_text=True, vertical="center")
                sh.column_dimensions[get_column_letter(i)].width = widths.get(h, 11)
            for r in sorted(rs, key=lambda r: (order.get(r["f"]["d"], 3), r["u"]["id"])):
                sh.append([srlib.spreadsheet_text(v) for v in line(r)])
                sh.cell(row=sh.max_row, column=2).fill = PatternFill("solid", fgColor=fills.get(r["f"]["d"], "FFFFFF"))
            sh.freeze_panes = "B2"
            sh.auto_filter.ref = sh.dimensions
        if pending_ids:
            sh = wb.create_sheet("Pending")
            sh.append(["ID", "Title", "Batch"])
            for rid in pending_ids:
                u = U.get(rid, {})
                sh.append([srlib.spreadsheet_text(v) for v in [rid, u.get("title", "[unknown record]"), u.get("batch", "")]])
        log_path = os.path.join(out, f"{P}_screening_log.xlsx")
        wb.save(log_path)
    except ImportError:
        log_path = os.path.join(out, f"{P}_screening_log_*.csv")
        with open(os.path.join(out, f"{P}_screening_log_Summary.csv"), "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f, quoting=csv.QUOTE_ALL).writerows([srlib.spreadsheet_text(v) for v in row] for row in summary)
        for name, rs in cols:
            with open(os.path.join(out, f"{P}_screening_log_{name}.csv"), "w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f, quoting=csv.QUOTE_ALL)
                w.writerow(head)
                w.writerows([srlib.spreadsheet_text(v) for v in line(r)] for r in
                            sorted(rs, key=lambda r: (order.get(r["f"]["d"], 3), r["u"]["id"])))
        with open(os.path.join(out, f"{P}_screening_log_Pending.csv"), "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, quoting=csv.QUOTE_ALL)
            w.writerow(["ID", "Title", "Batch"])
            for rid in pending_ids:
                u = U.get(rid, {})
                w.writerow([srlib.spreadsheet_text(v) for v in [rid, u.get("title", ""), u.get("batch", "")]])
        print("openpyxl not installed: wrote CSV files instead of .xlsx (pip install openpyxl for Excel)")

    # ------------------------------------------------------------------ RIS
    def raw_for(u):
        ms = sorted(u["members"], key=lambda i: srlib.db_rank((RAW[i].get("DB") or [""])[0]))
        return RAW[ms[0]]

    def extra(r):
        u, f = r["u"], r["f"]
        lab = f"{P}-{f['d'].capitalize()}" + (f"-{f['code']}" if f["d"] == "exclude" else "")
        note = (f"sr-screener {P} [{u['id']}]: {f['d'].upper()} ({labels.get(f['code'], f['code'])}) - {f.get('why', '')}"
                f" | A: {r['A'].get('d', '')}/{r['A'].get('code', '')} | B: {r['B'].get('d', '')}/{r['B'].get('code', '')}"
                f" | by {f['by']}")
        tags = [("LB", lab), ("ID", u["id"]), ("N1", note)]
        if a.tag_keywords:
            tags.append(("KW", lab))
        return tags

    ris_files = []
    for n, d in ((1, "include"), (2, "unclear"), (3, "exclude")):
        rs = [r for r in rows if r["f"]["d"] == d]
        p = os.path.join(out, f"{P}_{n}_{d}.ris")
        srlib.write_ris(p, [(raw_for(r["u"]), extra(r)) for r in rs])
        ris_files.append((os.path.basename(p), len(rs)))

    # ------------------------------------------------------------- PRISMA md
    md = [f"# PRISMA 2020 counts - {'title/abstract' if a.stage == 'ta' else 'full-text'} stage", ""]
    if not complete:
        md += [f"> **Provisional:** {len(pending_ids)} records are still pending.", ""]
        md += ["> " + problem for problem in state_problems]
        if a.stage == "ft" and ta_pending_ids:
            md += [f"> Title/abstract decisions or rechecks remain pending, including {len(ta_qc_ids)} "
                   "required title/abstract QC rechecks. These full-text counts are not final.", ""]
    if a.stage == "ta":
        dbs = "; ".join(f"{k} n = {v}" for k, v in counts["identified_by_database"].items())
        excl = "<br/>".join(f"{labels.get(c, c)}: {n}" for c, n in exc_sorted)
        md += ["```mermaid", "flowchart TD",
               f'  I["Records identified from databases (n = {counts["identified_total"]})<br/>{dbs}"]',
               f'  R["Records removed before screening:<br/>duplicate records (n = {counts["duplicates_removed"]})"]',
               f'  S["Records screened (n = {counts["records_screened"]})"]',
               f'  X["Records excluded (n = {counts["records_excluded"]})<br/>{excl}"]',
               f'  F["Reports sought for retrieval (n = {counts["reports_sought_for_retrieval"]})"]',
               "  I --> R", "  I --> S", "  S --> X", "  S --> F", "```", ""]
    else:
        reasons = "<br/>".join(f"{k}: {v}" for k, v in counts["excluded_by_reason"].items())
        md += ["```mermaid", "flowchart TD",
               f'  F["Reports sought for retrieval (n = {counts["reports_sought_for_retrieval"]})"]',
               f'  N["Reports not retrieved (n = {counts["reports_not_retrieved"]})"]',
               f'  E["Reports assessed for eligibility (n = {counts["reports_assessed"]})"]',
               f'  X["Reports excluded (n = {counts["reports_excluded"]})<br/>{reasons}"]',
               f'  C["Reports of included studies (n = {counts["studies_included_reports"]})"]',
               f'  U["Reports awaiting classification (n = {counts["reports_unclear_awaiting_classification"]})"]',
               "  F --> N", "  F --> E", "  E --> X", "  E --> C", "  E --> U", "```", ""]
    md += ["```json", json.dumps({k: v for k, v in counts.items() if k != "ta"}, indent=1, ensure_ascii=False), "```"]
    with open(os.path.join(out, f"{P}_prisma_counts.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(md) + "\n")

    # --------------------------------------------------------- methods text
    mpath = os.path.join(out, f"{P}_methods_selection.md")
    with open(mpath, "w", encoding="utf-8") as f:
        if not complete:
            f.write("# Methods text not generated\n\nScreening is incomplete "
                    f"({len(pending_ids)} records pending). Finish the run first; numbers must not be reported "
                    "from a partial run.\n")
            for problem in state_problems:
                f.write("\n" + problem + ".\n")
            if a.stage == "ft" and ta_pending_ids:
                f.write(f"\nThe title/abstract stage is unfinished, including {len(ta_qc_ids)} required "
                        "title/abstract QC rechecks. Decide these items before producing final full-text counts "
                        "or methods text.\n")
        else:
            po = agree.get("observed_agreement")
            po_txt = f"{po * 100:.1f}%" if isinstance(po, (int, float)) else "[n/a]"
            n_conf = agree.get("conflicts", 0)
            qc_n = sum(1 for r in rows if r["Q"])
            qc_adv = sum(1 for r in rows if r["pre_qc"])
            tool = (f"the sr-screener skill (version {srlib.VERSION}) of the open-source Academic Research Skills suite "
                    f"({REPO_URL}), running in Claude Code")
            access = ("The reviewer agents could only read the record file they were given (no web access). "
                      if (cfg.get("agent_type") or "").split(":")[-1] == "screening_reviewer_agent" else "")
            if a.stage == "ta":
                conflict = (f"Disagreements between advancing (include or unclear) and excluding a record "
                            f"({n_conf} records) were resolved by a third AI reviewer ({mlab('ADJ')}) that re-read the record. "
                            if cfg["conflict_policy"] == "adjudicate" else
                            f"Records advanced by either reviewer were retained for full-text assessment "
                            f"({n_conf} disagreements). ")
                qc = (f"Because records excluded by both reviewers are never adjudicated, a reproducible random sample "
                      f"of those joint exclusions, together with exclusions matching keyword signals for the core "
                      f"criteria ({qc_n} records in total), was re-screened by a senior AI reviewer ({mlab('QC')}) that "
                      f"did not see the earlier decisions; {qc_adv} were advanced as a result. " if qc_n else "")
                pc = srlib.load_json(os.path.join(work, "pilot_check.json")) or {}
                po_log = srlib.load_json(os.path.join(work, "pilot_override.json")) or []
                pilot_txt = ""
                if pc.get("compared"):
                    pilot_txt = (f"Before the full run, the AI screening was piloted on {pc['compared']} records that the "
                                 f"review team had labelled independently; the AI excluded "
                                 f"{len(pc.get('missed_advances', []))} of the {pc.get('human_advanced', 0)} records the "
                                 "team advanced. [TO COMPLETE: pilot rounds and protocol amendments.] ")
                if po_log:
                    pilot_txt += ("[TO COMPLETE: pilot gate or scope override(s) were recorded; "
                                  f"latest reason: {po_log[-1]['reason']}; describe their scope and revision.] ")
                seed_txt = ""
                if seeds:
                    ok = [s for s in seeds if s.get("id") and D.get(s["id"], {}).get("final", {}).get("d") in srlib.ADVANCE]
                    if len(ok) == len(seeds):
                        seed_txt = ("The known eligible study used as a sensitivity check was retrieved by the search and "
                                    "advanced. " if len(seeds) == 1 else
                                    f"All {len(seeds)} known eligible studies used as a sensitivity check were retrieved "
                                    "by the search and advanced. ")
                    else:
                        seed_txt = (f"{len(ok)} of {len(seeds)} known eligible studies used as a sensitivity check were "
                                    "retrieved and advanced [TO COMPLETE: explain the others]. ")
                dbs = ", ".join(f"{k} (n = {v})" for k, v in counts["identified_by_database"].items())
                f.write("# Selection process - title/abstract screening\n\n")
                f.write(
                    f"Records retrieved from {dbs} were combined and de-duplicated with {tool}, which matches PMIDs, "
                    f"DOIs and normalised titles with publication year (plus or minus one year); "
                    f"{counts['duplicates_removed']} duplicates were removed, leaving {counts['records_screened']} "
                    "unique records. Titles and abstracts were screened against the pre-specified eligibility criteria "
                    f"of the review protocol. Each record was assessed independently by two AI reviewers "
                    f"({pair('A', 'B')}) that were given the same written protocol and decision rules and did "
                    f"not see each other's decisions. {access}{conflict}The two reviewers agreed on advance-versus-exclude "
                    f"decisions for {po_txt} of records (Cohen's kappa = {agree.get('kappa')}; prevalence-adjusted "
                    f"bias-adjusted kappa = {agree.get('pabak')}). {pilot_txt}{qc}{seed_txt}In total, {counts['records_excluded']} "
                    f"records were excluded and {counts['reports_sought_for_retrieval']} records "
                    f"({counts['advanced_include']} judged eligible and {counts['advanced_unclear']} with insufficient "
                    "information) were advanced to full-text assessment. [TO COMPLETE - human verification: who checked "
                    "the AI decisions (for example all advanced records and a random sample of excluded records), how "
                    "disagreements with the AI were resolved, and the final human-confirmed numbers.]\n\n")
            else:
                reasons = "; ".join(f"{k} (n = {v})" for k, v in counts["excluded_by_reason"].items()) or "none"
                f.write("# Selection process - full-text assessment\n\n")
                f.write(
                    f"Of {counts['reports_sought_for_retrieval']} reports sought for retrieval, "
                    f"{counts['reports_not_retrieved']} could not be retrieved. {counts['reports_assessed']} full-text "
                    f"reports were assessed independently by two AI reviewers ({pair('FTA', 'FTB')}) using "
                    f"{tool}; each reviewer recorded one exclusion reason and the page supporting it. {access}"
                    f"Disagreements were resolved by a third AI reviewer ({mlab('FTADJ')}). The reviewers agreed on "
                    f"include-versus-exclude decisions for {po_txt} of reports (Cohen's kappa = {agree.get('kappa')}). "
                    f"{counts['reports_excluded']} reports were excluded ({reasons}); "
                    f"{counts['reports_unclear_awaiting_classification']} await classification; "
                    f"{counts['studies_included_reports']} reports met the eligibility criteria. [TO COMPLETE - human "
                    "verification of every included and excluded report, and how disagreements were resolved.]\n\n")
            used = sorted({mlab(k) for k in (("A", "B", "ADJ") if a.stage == "ta" else ("FTA", "FTB", "FTADJ"))})
            f.write("# Use of AI tools (draft statement)\n\n")
            f.write(
                f"{'Title/abstract screening' if a.stage == 'ta' else 'Full-text eligibility assessment'} was assisted "
                f"by large language models ({'; '.join(used)}) orchestrated by {tool}, between [TO COMPLETE: dates]. "
                "The AI reviewers applied the review's pre-specified eligibility criteria verbatim and judged only the "
                "text of each record or report; every decision and its reason is archived in the supplementary "
                "screening log. [TO COMPLETE: human oversight - which decisions the review team checked and how.] "
                "The review team takes full responsibility for the selection of studies. Check the target journal's "
                "AI policy and current guidance on AI in evidence synthesis (for example RAISE) before submission; "
                "PRISMA 2020 item 8 asks for details of any automation tools used in the selection process.\n")

    # --------------------------------------------------------------- corpus
    when = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    n_ok, rejected = write_corpus(os.path.join(out, f"{P}_literature_corpus.yaml"), adv_rows, a.stage, when)

    print(f"records with final decision: {len(rows)} / {len(universe)}" + ("" if complete else "  ** INCOMPLETE **"))
    print(f"include {count('include')}  unclear {count('unclear')}  exclude {count('exclude')}  by: {by}")
    print(f"log: {log_path}")
    for fn, n in ris_files:
        print(f"RIS: {fn} ({n})")
    print(f"PRISMA: {P}_prisma_counts.md / .json   methods: {os.path.basename(mpath)}")
    print(f"literature corpus: {n_ok} entries" + (f", {len(rejected)} left out (no authors/year)" if rejected else ""))


if __name__ == "__main__":
    main()
